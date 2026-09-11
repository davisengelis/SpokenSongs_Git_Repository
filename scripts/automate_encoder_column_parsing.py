from openpyxl import load_workbook
from pathlib import Path
import re

# --------------------------------------------------
# SETTINGS
# --------------------------------------------------

filename_list = Path("/Users/davisengelis/SpokenSongs_CodeBook/data/mei/ms-files/Temp_IevasFaili/filenames_no_extension.txt")

metadata_file = Path("/Users/davisengelis/SpokenSongs_CodeBook/metadata/LindaMetadata_6.xlsx")

output_file = Path("/Users/davisengelis/SpokenSongs_CodeBook/metadata/LindaMetadata_6_Ieva.xlsx")

encoder_name = "Ieva"


# --------------------------------------------------
# CONVERT FILENAME ID TO VitolinsSource FORMAT
# --------------------------------------------------

def filename_to_vitolins_source(filename_id):
    """
    Examples:

    V6-0001
        -> V VI, 1

    V3-0036-37
        -> V III, 36, 37

    V3-0032-35-36
        -> V III, 32, 35-36
    """

    match = re.fullmatch(r"V(\d+)-(\d+)(?:-(.+))?", filename_id)

    if not match:
        raise ValueError(f"Unrecognized filename format: {filename_id}")

    volume_number = int(match.group(1))
    song_number = int(match.group(2))
    extra = match.group(3)

    # Arabic -> Roman numeral
    roman_numbers = {
        1: "I",
        2: "II",
        3: "III",
        4: "IV",
        5: "V",
        6: "VI",
        7: "VII",
        8: "VIII",
        9: "IX",
        10: "X",
    }

    if volume_number not in roman_numbers:
        raise ValueError(
            f"No Roman numeral mapping defined for V{volume_number}"
        )

    source = f"V {roman_numbers[volume_number]}, {song_number}"

    if extra:
        extra_parts = extra.split("-")

        # One extra number:
        # V3-0036-37 -> V III, 36, 37
        if len(extra_parts) == 1:
            source += f", {int(extra_parts[0])}"

        # Two or more:
        # V3-0032-35-36 -> V III, 32, 35-36
        else:
            cleaned_extra = "-".join(
                str(int(part)) for part in extra_parts
            )
            source += f", {cleaned_extra}"

    return source


# --------------------------------------------------
# READ FILE LIST
# --------------------------------------------------

with open(filename_list, "r", encoding="utf-8") as f:
    filename_ids = [
        line.strip()
        for line in f
        if line.strip()
    ]


# Convert IDs into metadata-table notation

wanted_sources = {}

for filename_id in filename_ids:
    try:
        source = filename_to_vitolins_source(filename_id)
        wanted_sources[source] = filename_id

    except ValueError as e:
        print(f"WARNING: {e}")


print(f"\nIDs read from file: {len(filename_ids)}")
print(f"Valid IDs converted: {len(wanted_sources)}")


# --------------------------------------------------
# OPEN METADATA WORKBOOK
# --------------------------------------------------

wb = load_workbook(metadata_file)
ws = wb.active


# --------------------------------------------------
# FIND COLUMNS BY HEADER NAME
# --------------------------------------------------

headers = {}

for cell in ws[1]:
    if cell.value:
        headers[str(cell.value).strip()] = cell.column


if "VitolinsSource" not in headers:
    raise ValueError(
        'Column "VitolinsSource" was not found.'
    )

if "Encoder" not in headers:
    raise ValueError(
        'Column "Encoder" was not found.'
    )


source_col = headers["VitolinsSource"]
encoder_col = headers["Encoder"]

# --------------------------------------------------
# MATCH ROWS AND UPDATE ENCODER
# --------------------------------------------------

found = set()
updated_count = 0


for row in range(2, ws.max_row + 1):

    source_value = ws.cell(
        row=row,
        column=source_col
    ).value

    if source_value is None:
        continue

    source_value = str(source_value).strip()

    matched_source = None

    for wanted_source in wanted_sources:

        # Exact match:
        # V I, 312 == V I, 312
        #
        # OR metadata contains additional indexing:
        # V I, 312, 146
        # starts with
        # V I, 312,
        if (
            source_value == wanted_source
            or source_value.startswith(wanted_source + ",")
        ):
            matched_source = wanted_source
            break

    if matched_source is not None:

        found.add(matched_source)

        encoder_cell = ws.cell(
            row=row,
            column=encoder_col
        )

        current_value = encoder_cell.value

        # Empty Encoder cell
        if current_value is None or str(current_value).strip() == "":
            encoder_cell.value = encoder_name
            updated_count += 1

        else:
            current_value = str(current_value).strip()

            existing_encoders = [
                name.strip()
                for name in current_value.split(",")
            ]

            if encoder_name not in existing_encoders:
                encoder_cell.value = (
                    current_value + ", " + encoder_name
                )
                updated_count += 1

# --------------------------------------------------
# REPORT FILES THAT WERE NOT FOUND
# --------------------------------------------------

missing = set(wanted_sources) - found


print("\n----------------------------")
print("RESULT")
print("----------------------------")

print(f"Songs requested: {len(wanted_sources)}")
print(f"Songs found:     {len(found)}")
print(f"Rows updated:    {updated_count}")
print(f"Not found:       {len(missing)}")


if missing:
    print("\nNot found in metadata table:")

    for source in sorted(missing):
        original_id = wanted_sources[source]
        print(f"  {original_id}  ->  {source}")


# --------------------------------------------------
# SAVE AS NEW FILE
# --------------------------------------------------

wb.save(output_file)

print(f"\nSaved new metadata file as:")
print(output_file)