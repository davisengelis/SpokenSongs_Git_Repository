#!/usr/bin/env python3
from __future__ import annotations

import argparse
import copy
import csv
import logging
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from lxml import etree
from openpyxl import load_workbook

PROJECT_NS = "https://jvlma.lv/ns/spoken-songs"
PROJECT_PREFIX = "ss"
XML_NS = "http://www.w3.org/XML/1998/namespace"
SINGER_SUFFIXES = tuple("abcdefg")
FILE_NUMBER_RE = re.compile(r"V2-(\d+)", re.IGNORECASE)
FILE_IDENTITY_RE = re.compile(r"^V2-(\d+)(?:-([a-z]))?$", re.IGNORECASE)
SOURCE_NUMBER_RE = re.compile(r"\bV\s*II\s*,\s*0*(\d+)\s*,?", re.IGNORECASE)

MEI_COLUMNS = [
    "workTitle", "unitID", "VitolinsSource", "collector", "dateSource", "placeSource",
    "placeSourcePast", "Latitude", "Longitude", "SingerID", "SingerFull",
    "SingerSurname_a", "SingerName_a", "SingerSurname_b", "SingerName_b",
    "SingerSurname_c", "SingerName_c", "SingerSurname_d", "SingerName_d",
    "SingerSurname_e", "SingerName_e", "SingerSurname_f", "SingerName_f",
    "SingerSurname_g", "SingerName_g", "SingerAlternative_surname",
    "SingerAlternative_name", "SingerBirthYear", "Singer_a_BirthYear",
    "Singer_b_BirthYear", "Singer_c_BirthYear", "Singer_d_BirthYear",
    "Singer_e_BirthYear", "Singer_f_BirthYear", "Singer_g_BirthYear",
    "SingerBirthPlace", "SingerLatitude", "SingerLongitude", "SingerMisc",
    "textLanguage", "witnessId",
]


class EnrichmentError(RuntimeError):
    pass


@dataclass(frozen=True)
class MetadataRow:
    excel_row: int
    source_number: int
    values: dict[str, str]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    text = unicodedata.normalize("NFC", str(value))
    return re.sub(r"\s+", " ", text).strip()


def canonical_header(value: Any) -> str:
    return clean_text(value).casefold()


def mei_qname(namespace: str, local_name: str) -> etree.QName:
    return etree.QName(namespace, local_name)


def project_qname(local_name: str) -> etree.QName:
    return etree.QName(PROJECT_NS, local_name)


def detect_mei_namespace(root: etree._Element) -> str:
    qname = etree.QName(root)
    if qname.localname != "mei":
        raise EnrichmentError(f"Root element is <{qname.localname}>, not <mei>.")
    if not qname.namespace:
        raise EnrichmentError("The MEI document has no XML namespace.")
    return qname.namespace


def read_metadata(metadata_path: Path, sheet_name: str | None) -> tuple[dict[int, list[MetadataRow]], list[str]]:
    workbook = load_workbook(metadata_path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise EnrichmentError(
                f"Worksheet {sheet_name!r} not found. Available: {', '.join(workbook.sheetnames)}"
            )
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active

    raw_headers = [clean_text(cell.value) for cell in worksheet[1]]
    header_lookup: dict[str, int] = {}
    duplicate_headers: list[str] = []
    for index, header in enumerate(raw_headers):
        if not header:
            continue
        key = canonical_header(header)
        if key in header_lookup:
            duplicate_headers.append(header)
        else:
            header_lookup[key] = index

      # Only complain if duplicated headers are actually used by this script.
    columns_used = {
        canonical_header("VitolinsSource"),
        canonical_header("meiStatus"),
    }

    columns_used.update(
        canonical_header(column)
        for column in MEI_COLUMNS
    )

    relevant_duplicates = []

    for header in duplicate_headers:
        if canonical_header(header) in columns_used:
            relevant_duplicates.append(header)

    if relevant_duplicates:
        raise EnrichmentError(
            "Duplicate metadata headers used by the script: "
            + ", ".join(repr(header) for header in relevant_duplicates)
        )

    required = ["VitolinsSource", "meiStatus"]
    missing_required = [
        column
        for column in required
        if canonical_header(column) not in header_lookup
    ]

    if missing_required:
        raise EnrichmentError(
            "Missing required metadata columns: "
            + ", ".join(missing_required)
        )

    available_columns = [
        column
        for column in MEI_COLUMNS
        if canonical_header(column) in header_lookup
    ]

    missing_optional = [
        column
        for column in MEI_COLUMNS
        if canonical_header(column) not in header_lookup
    ]

    rows_by_source: dict[int, list[MetadataRow]] = defaultdict(list)

    for excel_row, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        status = clean_text(
            row[header_lookup[canonical_header("meiStatus")]]
        )

        if status.casefold() != "r":
            continue

        source_text = clean_text(
            row[header_lookup[canonical_header("VitolinsSource")]]
        )

        match = SOURCE_NUMBER_RE.search(source_text)

        if not match:
            continue

        source_number = int(match.group(1))
        values: dict[str, str] = {}

        for column in available_columns:
            value = clean_text(
                row[header_lookup[canonical_header(column)]]
            )

            if value:
                values[column] = value

        rows_by_source[source_number].append(
            MetadataRow(
                excel_row=excel_row,
                source_number=source_number,
                values=values,
            )
        )

    workbook.close()
    return dict(rows_by_source), missing_optional


def extract_file_identity(path: Path) -> tuple[int, str | None]:
    match = FILE_IDENTITY_RE.fullmatch(path.stem)
    if not match:
        raise EnrichmentError(
            "Filename must follow V2-0002.mei or V2-0002-a.mei naming."
        )
    source_number = int(match.group(1))
    variant = match.group(2).casefold() if match.group(2) else None
    return source_number, variant


def extract_file_number(path: Path) -> int:
    source_number, _ = extract_file_identity(path)
    return source_number


def choose_metadata_row(source_number: int, candidates: list[MetadataRow] | None) -> MetadataRow:
    if not candidates:
        raise EnrichmentError(
            f"No metadata row with meiStatus='r' and VitolinsSource matching V II, {source_number}, was found."
        )
    if len(candidates) == 1:
        return candidates[0]

    first = candidates[0].values
    if all(candidate.values == first for candidate in candidates[1:]):
        logging.warning(
            "Source %s has %s identical eligible metadata rows (%s); using the first.",
            source_number,
            len(candidates),
            ", ".join(str(candidate.excel_row) for candidate in candidates),
        )
        return candidates[0]

    rows = ", ".join(str(candidate.excel_row) for candidate in candidates)
    raise EnrichmentError(
        f"Ambiguous metadata: source {source_number} matches different meiStatus='r' rows: {rows}."
    )


def add_text_element(parent: etree._Element, namespace: str, local_name: str, text: str, **attributes: str) -> etree._Element:
    element = etree.SubElement(parent, mei_qname(namespace, local_name))
    for key, value in attributes.items():
        if value:
            element.set(key, value)
    element.text = text
    return element


def append_standard_header_metadata(
    mei_head: etree._Element,
    mei_ns: str,
    metadata: dict[str, str],
    source_number: int,
    input_filename: str,
) -> None:
    """Add MEI-native descriptive metadata in schema-friendly header order."""
    title = metadata.get("workTitle") or f"Spoken song V II, {source_number}"
    source_id = metadata.get("VitolinsSource", f"V II, {source_number},")

    # 1. fileDesc — description of this digital MEI file and its immediate source.
    file_desc = etree.SubElement(mei_head, mei_qname(mei_ns, "fileDesc"))

    title_stmt = etree.SubElement(file_desc, mei_qname(mei_ns, "titleStmt"))
    add_text_element(title_stmt, mei_ns, "title", title)

    resp_stmt = etree.SubElement(title_stmt, mei_qname(mei_ns, "respStmt"))
    add_text_element(resp_stmt, mei_ns, "resp", "Metadata enrichment and MEI preparation")
    add_text_element(resp_stmt, mei_ns, "corpName", "Spoken Songs project")

    pub_stmt = etree.SubElement(file_desc, mei_qname(mei_ns, "pubStmt"))
    pub_resp = etree.SubElement(pub_stmt, mei_qname(mei_ns, "respStmt"))
    add_text_element(pub_resp, mei_ns, "corpName", "Spoken Songs project")
    add_text_element(pub_resp, mei_ns, "resp", "Responsible for this research encoding")

    source_desc = etree.SubElement(file_desc, mei_qname(mei_ns, "sourceDesc"))
    source = etree.SubElement(source_desc, mei_qname(mei_ns, "source"))
    source.set(etree.QName(XML_NS, "id"), f"source_V2_{source_number:04d}")

    bibl = etree.SubElement(source, mei_qname(mei_ns, "bibl"))
    add_text_element(
        bibl,
        mei_ns,
        "identifier",
        source_id,
        type="VitolinsSource",
    )

    if metadata.get("workTitle"):
        add_text_element(bibl, mei_ns, "title", metadata["workTitle"])

    if metadata.get("collector"):
        resp = etree.SubElement(bibl, mei_qname(mei_ns, "respStmt"))
        add_text_element(resp, mei_ns, "resp", "Collector")
        collector_name = add_text_element(
            resp,
            mei_ns,
            "persName",
            metadata["collector"],
        )
        collector_name.set("role", "collector")

    if metadata.get("dateSource"):
        add_text_element(
            bibl,
            mei_ns,
            "date",
            metadata["dateSource"],
            type="collection",
        )

    # 2. encodingDesc — how this machine-readable file was produced.
    encoding_desc = etree.SubElement(mei_head, mei_qname(mei_ns, "encodingDesc"))
    app_info = etree.SubElement(encoding_desc, mei_qname(mei_ns, "appInfo"))
    application = etree.SubElement(app_info, mei_qname(mei_ns, "application"))
    application.set(etree.QName(XML_NS, "id"), "spoken_songs_metadata_enricher")
    add_text_element(application, mei_ns, "name", "Spoken Songs metadata enricher")
    add_text_element(
        application,
        mei_ns,
        "p",
        "Preserved the MuseScore-exported musical content and generated metadata "
        "from LindaMetadata_1.xlsx.",
    )

    # 3. workList — MEI-native description of the song/work represented.
    work_list = etree.SubElement(mei_head, mei_qname(mei_ns, "workList"))
    work = etree.SubElement(work_list, mei_qname(mei_ns, "work"))
    work.set(etree.QName(XML_NS, "id"), f"work_V2_{source_number:04d}")

    add_text_element(
        work,
        mei_ns,
        "identifier",
        source_id,
        type="VitolinsSource",
    )
    add_text_element(work, mei_ns, "title", title)

    # Collector as a work-level contributor.
    if metadata.get("collector"):
        contributor = etree.SubElement(work, mei_qname(mei_ns, "contributor"))
        contributor.set("role", "collector")
        contributor.text = metadata["collector"]

    # Singers/performers as structured contributors with stable IDs where available.
    singer_ids = [
        clean_text(value)
        for value in re.split(r"\s*,\s*", metadata.get("SingerID", ""))
        if clean_text(value)
    ]

    singer_full_values = [
        clean_text(value)
        for value in re.split(r"\s*,\s*", metadata.get("SingerFull", ""))
        if clean_text(value)
    ]

    for index, suffix in enumerate(SINGER_SUFFIXES):
        surname = metadata.get(f"SingerSurname_{suffix}", "")
        given = metadata.get(f"SingerName_{suffix}", "")
        birth_year = metadata.get(f"Singer_{suffix}_BirthYear", "")

        if not surname and not given:
            continue

        display_name = " ".join(part for part in [given, surname] if part).strip()
        if not display_name and index < len(singer_full_values):
            display_name = singer_full_values[index]
        if not display_name:
            continue

        contributor = etree.SubElement(work, mei_qname(mei_ns, "contributor"))
        contributor.set("role", "performer")

        pers_name = etree.SubElement(contributor, mei_qname(mei_ns, "persName"))
        pers_name.text = display_name

        if index < len(singer_ids):
            singer_id = singer_ids[index]
            if re.fullmatch(r"[A-Za-z_][A-Za-z0-9._-]*", singer_id):
                pers_name.set(etree.QName(XML_NS, "id"), singer_id)

        if birth_year:
            pers_name.set("startdate", birth_year)

    # Collection circumstances belong to work history.
    if metadata.get("dateSource") or metadata.get("placeSource") or metadata.get("placeSourcePast"):
        history = etree.SubElement(work, mei_qname(mei_ns, "history"))
        event_list = etree.SubElement(history, mei_qname(mei_ns, "eventList"))
        event_list.set("type", "collection")
        event = etree.SubElement(event_list, mei_qname(mei_ns, "event"))

        if metadata.get("dateSource"):
            add_text_element(
                event,
                mei_ns,
                "date",
                metadata["dateSource"],
            )

        if metadata.get("placeSource"):
            add_text_element(
                event,
                mei_ns,
                "geogName",
                metadata["placeSource"],
                type="collectionPlace",
            )

        if metadata.get("placeSourcePast"):
            add_text_element(
                event,
                mei_ns,
                "geogName",
                metadata["placeSourcePast"],
                type="historicalName",
            )

    # Language of the encoded sung text.
    if metadata.get("textLanguage"):
        lang_usage = etree.SubElement(work, mei_qname(mei_ns, "langUsage"))
        add_text_element(
            lang_usage,
            mei_ns,
            "language",
            metadata["textLanguage"],
        )


def append_revision_metadata(
    mei_head: etree._Element,
    mei_ns: str,
    source_number: int,
    input_filename: str,
) -> None:
    """Add revisionDesc last, as required by the MEI header content model."""
    revision_desc = etree.SubElement(mei_head, mei_qname(mei_ns, "revisionDesc"))
    change = etree.SubElement(revision_desc, mei_qname(mei_ns, "change"))
    change.set("isodate", date.today().isoformat())

    change_desc = etree.SubElement(change, mei_qname(mei_ns, "changeDesc"))
    add_text_element(
        change_desc,
        mei_ns,
        "p",
        f"Metadata injected for source V II, {source_number}; "
        f"musical content preserved from {input_filename}.",
    )

def append_project_metadata(
    mei_head: etree._Element,
    mei_ns: str,
    metadata: dict[str, str],
    source_number: int,
    excel_row: int,
) -> None:
    ext_meta = etree.SubElement(
        mei_head,
        mei_qname(mei_ns, "extMeta"),
        nsmap={PROJECT_PREFIX: PROJECT_NS},
    )
    record = etree.SubElement(ext_meta, project_qname("record"))
    record.set("sourceNumber", str(source_number))
    record.set("metadataExcelRow", str(excel_row))

    simple_fields = [
        "workTitle", "VitolinsSource", "collector", "dateSource", "placeSource",
        "placeSourcePast", "Latitude", "Longitude", "SingerID", "SingerFull",
        "SingerAlternative_surname", "SingerAlternative_name", "SingerBirthYear",
        "SingerBirthPlace", "SingerLatitude", "SingerLongitude", "SingerMisc",
        "textLanguage", "witnessId",
    ]
    for field in simple_fields:
        value = metadata.get(field)
        if value:
            child = etree.SubElement(record, project_qname(field))
            child.text = value

    singer_ids = [clean_text(v) for v in re.split(r"\s*,\s*", metadata.get("SingerID", "")) if clean_text(v)]
    singers = etree.SubElement(record, project_qname("singers"))
    count = 0
    for index, suffix in enumerate(SINGER_SUFFIXES):
        surname = metadata.get(f"SingerSurname_{suffix}", "")
        given = metadata.get(f"SingerName_{suffix}", "")
        birth_year = metadata.get(f"Singer_{suffix}_BirthYear", "")
        if not surname and not given and not birth_year:
            continue
        singer = etree.SubElement(singers, project_qname("singer"))
        singer.set("slot", suffix)
        if index < len(singer_ids):
            singer.set("ref", singer_ids[index])
        if given:
            etree.SubElement(singer, project_qname("givenName")).text = given
        if surname:
            etree.SubElement(singer, project_qname("surname")).text = surname
        if birth_year:
            etree.SubElement(singer, project_qname("birthYear")).text = birth_year
        count += 1
    if count == 0:
        record.remove(singers)



def append_witness_sources(
    mei_head: etree._Element,
    mei_ns: str,
    source_number: int,
    witnesses: list[tuple[str, str]],
) -> None:
    """Declare encoded variant witnesses in fileDesc/sourceDesc."""
    file_desc = mei_head.find(f"{{{mei_ns}}}fileDesc")
    if file_desc is None:
        raise EnrichmentError("Generated meiHead has no fileDesc.")
    source_desc = file_desc.find(f"{{{mei_ns}}}sourceDesc")
    if source_desc is None:
        raise EnrichmentError("Generated fileDesc has no sourceDesc.")

    for letter, filename in witnesses:
        source = etree.SubElement(source_desc, mei_qname(mei_ns, "source"))
        source.set(etree.QName(XML_NS, "id"), f"wit_V2_{source_number:04d}_{letter}")
        source.set("type", "witness")
        add_text_element(source, mei_ns, "identifier", filename, type="MEIWitness")
        annot = etree.SubElement(source, mei_qname(mei_ns, "annot"))
        annot.text = f"Variant witness {letter}."


def collect_xml_ids(root: etree._Element) -> set[str]:
    return {
        value
        for value in root.xpath("//@xml:id", namespaces={"xml": XML_NS})
        if value
    }


def prefix_ids_and_internal_refs(
    element: etree._Element,
    prefix: str,
    witness_ids: set[str],
) -> None:
    """Make copied witness IDs unique and repair internal fragment references."""
    xml_id_key = etree.QName(XML_NS, "id")
    reference_attributes = {
        "startid", "endid", "sameas", "copyof", "corresp",
        "next", "prev", "plist", "target", "facs", "source",
    }

    for node in element.iter():
        old_id = node.get(xml_id_key)
        if old_id:
            node.set(xml_id_key, f"{prefix}_{old_id}")

        for attr_key, attr_value in list(node.attrib.items()):
            local_name = etree.QName(attr_key).localname if attr_key.startswith("{") else attr_key
            if local_name not in reference_attributes or not attr_value:
                continue

            tokens = attr_value.split()
            changed = False
            repaired: list[str] = []
            for token in tokens:
                if token.startswith("#") and token[1:] in witness_ids:
                    repaired.append(f"#{prefix}_{token[1:]}")
                    changed = True
                else:
                    repaired.append(token)
            if changed:
                node.set(attr_key, " ".join(repaired))


def measure_map_for_tree(
    tree: etree._ElementTree,
    mei_ns: str,
    filename: str,
) -> tuple[etree._Element, list[str], dict[str, etree._Element]]:
    sections = tree.xpath(".//m:section", namespaces={"m": mei_ns})
    if len(sections) != 1:
        raise EnrichmentError(
            f"{filename}: expected exactly one <section>, found {len(sections)}."
        )
    section = sections[0]
    measures = section.xpath("./m:measure", namespaces={"m": mei_ns})
    numbers = [clean_text(measure.get("n")) for measure in measures]
    if not measures:
        raise EnrichmentError(f"{filename}: no direct child <measure> elements found.")
    if any(not number for number in numbers):
        raise EnrichmentError(f"{filename}: every measure must have @n for variant alignment.")
    if len(set(numbers)) != len(numbers):
        raise EnrichmentError(f"{filename}: duplicate measure numbers found: {numbers}.")
    return section, numbers, {measure.get("n"): measure for measure in measures}


def scoredef_signature(tree: etree._ElementTree, mei_ns: str) -> list[bytes]:
    """Compare notation-defining scoreDef structures while ignoring formatting whitespace."""
    signatures: list[bytes] = []
    for scoredef in tree.xpath(".//m:scoreDef", namespaces={"m": mei_ns}):
        node = copy.deepcopy(scoredef)
        for child in node.iter():
            if child.text is not None and not child.text.strip():
                child.text = None
            if child.tail is not None and not child.tail.strip():
                child.tail = None
        signatures.append(etree.tostring(node, method="c14n"))
    return signatures


def validate_variant_letters(letters: list[str]) -> list[str]:
    if len(letters) < 2:
        raise EnrichmentError(
            "Metadata says witnessId='IR', but fewer than two variant files were found."
        )
    if len(set(letters)) != len(letters):
        raise EnrichmentError("Duplicate variant letters were found.")
    highest = max(ord(letter) for letter in letters)
    expected = [chr(code) for code in range(ord("a"), highest + 1)]
    missing = [letter for letter in expected if letter not in letters]
    if missing:
        raise EnrichmentError(
            "Variant sequence is not contiguous; missing: " + ", ".join(missing)
        )
    if letters[0] != "a":
        raise EnrichmentError("Variant sequence must begin with '-a'.")
    return missing


def collate_variant_files(
    variant_paths: list[Path],
    output_path: Path,
    metadata_row: MetadataRow,
) -> dict[str, int]:
    """Merge separate witness MEIs into one measure-level <app>/<rdg> document."""
    parsed: list[tuple[str, Path, etree._ElementTree, str]] = []
    for path in sorted(variant_paths, key=lambda p: extract_file_identity(p)[1] or ""):
        source_number, letter = extract_file_identity(path)
        if source_number != metadata_row.source_number:
            raise EnrichmentError(f"Unexpected source number in {path.name}.")
        if letter is None:
            raise EnrichmentError(
                f"{path.name}: IR records must use variant filenames such as V2-0001-a.mei."
            )
        parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False, no_network=True)
        tree = etree.parse(str(path), parser)
        mei_ns = detect_mei_namespace(tree.getroot())
        parsed.append((letter, path, tree, mei_ns))

    letters = [item[0] for item in parsed]
    validate_variant_letters(letters)

    namespaces = {item[3] for item in parsed}
    if len(namespaces) != 1:
        raise EnrichmentError("Variant witnesses use different MEI namespaces.")
    mei_ns = parsed[0][3]

    versions = {item[2].getroot().get("meiversion", "") for item in parsed}
    if len(versions) != 1:
        raise EnrichmentError(
            "Variant witnesses use different MEI versions: " + ", ".join(sorted(versions))
        )

    # The base witness supplies shared non-measure structure (scoreDef, page/system breaks, etc.).
    base_letter, base_path, base_tree, _ = parsed[0]
    base_section, base_numbers, _ = measure_map_for_tree(base_tree, mei_ns, base_path.name)
    base_scoredefs = scoredef_signature(base_tree, mei_ns)

    measure_maps: dict[str, dict[str, etree._Element]] = {}
    witness_ids: dict[str, set[str]] = {}
    measure_counts: dict[str, int] = {}

    for letter, path, tree, _ in parsed:
        _, numbers, mapping = measure_map_for_tree(tree, mei_ns, path.name)
        measure_counts[letter] = len(numbers)
        if numbers != base_numbers:
            raise EnrichmentError(
                f"Measure-number mismatch: {base_path.name} has {base_numbers}; "
                f"{path.name} has {numbers}."
            )
        if scoredef_signature(tree, mei_ns) != base_scoredefs:
            raise EnrichmentError(
                f"{path.name}: scoreDef structure differs from witness {base_letter}."
            )
        measure_maps[letter] = mapping
        witness_ids[letter] = collect_xml_ids(tree.getroot())

    # Replace each base-witness measure with an apparatus containing every reading.
    for base_measure in list(base_section.xpath("./m:measure", namespaces={"m": mei_ns})):
        number = base_measure.get("n")
        app = etree.Element(mei_qname(mei_ns, "app"))
        safe_number = re.sub(r"[^A-Za-z0-9_.-]+", "_", number or "")
        app.set(etree.QName(XML_NS, "id"), f"app_m{safe_number}")

        for letter, path, _, _ in parsed:
            rdg = etree.SubElement(app, mei_qname(mei_ns, "rdg"))
            rdg.set("source", f"#wit_V2_{metadata_row.source_number:04d}_{letter}")
            rdg.set(etree.QName(XML_NS, "id"), f"rdg_m{safe_number}_{letter}")

            measure_copy = copy.deepcopy(measure_maps[letter][number])
            prefix_ids_and_internal_refs(
                measure_copy,
                f"w{letter}",
                witness_ids[letter],
            )
            rdg.append(measure_copy)

        base_section.replace(base_measure, app)

    root = base_tree.getroot()
    old_head = root.find(f"{{{mei_ns}}}meiHead")
    witness_list = [(letter, path.name) for letter, path, _, _ in parsed]
    filenames = ", ".join(path.name for _, path, _, _ in parsed)
    new_head = create_generated_mei_head(
        old_head,
        mei_ns,
        metadata_row,
        filenames,
        witnesses=witness_list,
    )

    if old_head is not None:
        index = root.index(old_head)
        root.remove(old_head)
        root.insert(index, new_head)
    else:
        root.insert(0, new_head)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    etree.indent(base_tree, space="  ")
    base_tree.write(
        str(output_path),
        encoding="UTF-8",
        xml_declaration=True,
        pretty_print=True,
    )
    return measure_counts


def write_diagnostics_csv(path: Path, rows: list[dict[str, str]]) -> None:
    fields = [
        "source_number", "vitolins_source", "excel_row", "witnessId", "mode",
        "variants_found", "variant_count", "measure_counts", "output_file",
        "status", "message",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def create_generated_mei_head(
    old_head: etree._Element | None,
    mei_ns: str,
    metadata_row: MetadataRow,
    input_filename: str,
    witnesses: list[tuple[str, str]] | None = None,
) -> etree._Element:
    mei_head = etree.Element(
        mei_qname(mei_ns, "meiHead"),
        nsmap={None: mei_ns, PROJECT_PREFIX: PROJECT_NS},
    )
    mei_head.set("type", "music")
    old_copy = copy.deepcopy(old_head) if old_head is not None else None

    # MEI header order:
    # fileDesc -> encodingDesc -> workList -> extMeta -> revisionDesc
    append_standard_header_metadata(
        mei_head,
        mei_ns,
        metadata_row.values,
        metadata_row.source_number,
        input_filename,
    )

    if witnesses:
        append_witness_sources(
            mei_head,
            mei_ns,
            metadata_row.source_number,
            witnesses,
        )

    # Lossless project-specific metadata.
    append_project_metadata(
        mei_head,
        mei_ns,
        metadata_row.values,
        metadata_row.source_number,
        metadata_row.excel_row,
    )

    # Preserve the original MuseScore-generated header for auditability.
    if old_copy is not None:
        ext_meta = etree.SubElement(
            mei_head,
            mei_qname(mei_ns, "extMeta"),
            nsmap={PROJECT_PREFIX: PROJECT_NS},
        )
        previous = etree.SubElement(ext_meta, project_qname("previousMeiHead"))
        previous.append(old_copy)

    # revisionDesc must come after extMeta.
    append_revision_metadata(
        mei_head,
        mei_ns,
        metadata_row.source_number,
        input_filename,
    )

    return mei_head

def enrich_file(input_path: Path, output_path: Path, metadata_row: MetadataRow) -> None:
    parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False, no_network=True)
    tree = etree.parse(str(input_path), parser)
    root = tree.getroot()
    mei_ns = detect_mei_namespace(root)
    old_head = root.find(f"{{{mei_ns}}}meiHead")
    new_head = create_generated_mei_head(old_head, mei_ns, metadata_row, input_path.name)
    if old_head is not None:
        index = root.index(old_head)
        root.remove(old_head)
        root.insert(index, new_head)
    else:
        root.insert(0, new_head)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Re-indent the complete XML tree after inserting the generated metadata.
    # This makes <meiHead> and the rest of the MEI much easier to inspect
    # in text editors such as VS Code.
    etree.indent(tree, space="  ")

    tree.write(
        str(output_path),
        encoding="UTF-8",
        xml_declaration=True,
        pretty_print=True,
    )


def iter_mei_files(input_dir: Path, recursive: bool) -> Iterable[Path]:
    pattern = "**/*" if recursive else "*"
    yield from sorted(
        path for path in input_dir.glob(pattern)
        if path.is_file() and path.suffix.casefold() == ".mei"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Enrich raw MuseScore MEI files with Excel metadata.")
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--metadata", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--sheet", default=None, help="Excel worksheet name; active sheet by default.")
    parser.add_argument("--recursive", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--diagnostics-csv",
        type=Path,
        default=None,
        help="Diagnostics CSV path; defaults to OUTPUT_DIR/enrichment_diagnostics.csv.",
    )
    parser.add_argument("--log-level", choices=["DEBUG", "INFO", "WARNING", "ERROR"], default="INFO")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s: %(message)s")

    if not args.input_dir.is_dir():
        logging.error("Input directory does not exist: %s", args.input_dir)
        return 2
    if not args.metadata.is_file():
        logging.error("Metadata workbook does not exist: %s", args.metadata)
        return 2

    try:
        rows_by_source, missing_optional = read_metadata(args.metadata, args.sheet)
    except Exception as exc:
        logging.error("Could not read metadata: %s", exc)
        return 2

    if missing_optional:
        logging.warning("Optional metadata columns not found and skipped: %s", ", ".join(missing_optional))

    input_files = list(iter_mei_files(args.input_dir, args.recursive))
    if not input_files:
        logging.error("No .mei files found in %s", args.input_dir)
        return 2

    # Group unsuffixed songs and -a/-b/-c witnesses by their V2 source number.
    grouped: dict[int, list[Path]] = defaultdict(list)
    filename_failures: list[tuple[Path, str]] = []
    for input_path in input_files:
        try:
            source_number, _ = extract_file_identity(input_path)
            grouped[source_number].append(input_path)
        except Exception as exc:
            filename_failures.append((input_path, str(exc)))
            logging.error("%s: %s", input_path.name, exc)

    processed = 0
    variant_groups = 0
    failures: list[tuple[str, str]] = [
        (path.name, message) for path, message in filename_failures
    ]
    diagnostics: list[dict[str, str]] = []

    for source_number in sorted(grouped):
        group_paths = sorted(grouped[source_number])
        metadata_row: MetadataRow | None = None
        output_path = args.output_dir / f"V2-{source_number:04d}.mei"

        try:
            metadata_row = choose_metadata_row(source_number, rows_by_source.get(source_number))
            witness_id = clean_text(metadata_row.values.get("witnessId", ""))
            is_variants = witness_id.casefold() == "ir"
            identities = [extract_file_identity(path) for path in group_paths]
            letters = sorted(letter for _, letter in identities if letter is not None)
            unsuffixed = [path for path, (_, letter) in zip(group_paths, identities) if letter is None]

            if output_path.exists() and not args.overwrite:
                raise EnrichmentError(
                    f"Output already exists: {output_path}. Use --overwrite to replace it."
                )

            if is_variants:
                if unsuffixed:
                    raise EnrichmentError(
                        "Metadata says witnessId='IR', but an unsuffixed raw MEI is present: "
                        + ", ".join(path.name for path in unsuffixed)
                    )
                validate_variant_letters(letters)
                measure_counts = collate_variant_files(group_paths, output_path, metadata_row)
                variant_groups += 1
                processed += 1
                message = "Merged variant witnesses successfully."
                logging.info(
                    "V2-%04d variants %s -> %s",
                    source_number,
                    ",".join(letters),
                    output_path,
                )
                diagnostics.append({
                    "source_number": str(source_number),
                    "vitolins_source": metadata_row.values.get("VitolinsSource", ""),
                    "excel_row": str(metadata_row.excel_row),
                    "witnessId": witness_id,
                    "mode": "variants",
                    "variants_found": ";".join(letters),
                    "variant_count": str(len(letters)),
                    "measure_counts": ";".join(
                        f"{letter}:{measure_counts[letter]}" for letter in letters
                    ),
                    "output_file": output_path.name,
                    "status": "success",
                    "message": message,
                })
            else:
                if letters:
                    raise EnrichmentError(
                        "Variant-style filename(s) found, but metadata witnessId is not 'IR': "
                        + ", ".join(path.name for path in group_paths)
                    )
                if len(unsuffixed) != 1:
                    raise EnrichmentError(
                        f"Expected one raw MEI for a non-IR song; found {len(unsuffixed)}."
                    )
                enrich_file(unsuffixed[0], output_path, metadata_row)
                processed += 1
                logging.info(
                    "%s -> %s (source %s, Excel row %s)",
                    unsuffixed[0].name,
                    output_path,
                    source_number,
                    metadata_row.excel_row,
                )
                diagnostics.append({
                    "source_number": str(source_number),
                    "vitolins_source": metadata_row.values.get("VitolinsSource", ""),
                    "excel_row": str(metadata_row.excel_row),
                    "witnessId": witness_id,
                    "mode": "single",
                    "variants_found": "",
                    "variant_count": "0",
                    "measure_counts": "",
                    "output_file": output_path.name,
                    "status": "success",
                    "message": "Enriched single MEI successfully.",
                })

        except Exception as exc:
            label = f"V2-{source_number:04d}"
            failures.append((label, str(exc)))
            logging.error("%s: %s", label, exc)

            witness_id = ""
            vitolins_source = ""
            excel_row = ""
            if metadata_row is not None:
                witness_id = metadata_row.values.get("witnessId", "")
                vitolins_source = metadata_row.values.get("VitolinsSource", "")
                excel_row = str(metadata_row.excel_row)

            try:
                letters = sorted(
                    letter
                    for _, letter in (extract_file_identity(path) for path in group_paths)
                    if letter is not None
                )
            except Exception:
                letters = []

            diagnostics.append({
                "source_number": str(source_number),
                "vitolins_source": vitolins_source,
                "excel_row": excel_row,
                "witnessId": witness_id,
                "mode": "variants" if clean_text(witness_id).casefold() == "ir" else "single",
                "variants_found": ";".join(letters),
                "variant_count": str(len(letters)),
                "measure_counts": "",
                "output_file": output_path.name,
                "status": "error",
                "message": str(exc),
            })

    diagnostics_path = args.diagnostics_csv or (args.output_dir / "enrichment_diagnostics.csv")
    for path, message in filename_failures:
        diagnostics.append({
            "source_number": "",
            "vitolins_source": "",
            "excel_row": "",
            "witnessId": "",
            "mode": "unknown",
            "variants_found": "",
            "variant_count": "",
            "measure_counts": "",
            "output_file": "",
            "status": "error",
            "message": f"{path.name}: {message}",
        })
    write_diagnostics_csv(diagnostics_path, diagnostics)

    print(f"\nRaw MEI files found:    {len(input_files)}")
    print(f"Song groups found:      {len(grouped)}")
    print(f"Variant groups merged:  {variant_groups}")
    print(f"Successfully produced:  {processed}")
    print(f"Failed or ambiguous:    {len(failures)}")
    print(f"Output directory:       {args.output_dir}")
    print(f"Diagnostics CSV:        {diagnostics_path}")

    if missing_optional:
        print(f"\nOptional metadata columns absent: {len(missing_optional)}")
        for column in missing_optional:
            print(f"  - {column}")
    else:
        print("\nOptional metadata columns absent: none")

    if failures:
        print("\nFailures:")
        for label, message in failures:
            print(f"  - {label}: {message}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
