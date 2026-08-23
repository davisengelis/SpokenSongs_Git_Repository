from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = "/Users/davisengelis/SpokenSongs_CodeBook/metadata/LindaMetadata_1.xlsx"

wb = load_workbook(path, read_only=True, data_only=True)
ws = wb.active

headers = defaultdict(list)

for column_number, cell in enumerate(ws[1], start=1):
    raw = cell.value
    normalized = "" if raw is None else " ".join(str(raw).split()).casefold()

    if normalized:
        headers[normalized].append(
            (get_column_letter(column_number), raw)
        )

for normalized, occurrences in headers.items():
    if len(occurrences) > 1:
        print(f"Duplicate normalized header: {normalized!r}")
        for column, raw in occurrences:
            print(f"  Column {column}: {raw!r}")